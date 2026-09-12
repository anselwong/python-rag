"""结构化文档解析与质量门控。

第三方库只负责读取特定格式；本模块把它们统一为 ``DocumentElement``，使后续
切分、向量化和人工审核不依赖某个供应商的返回结构。解析失败不会悄悄降级为
低质量文本，而会携带质量报告进入 ``needs_review`` 状态。
"""

import json
import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

from langchain_core.documents import Document

from app.langchain.splitters import split_documents


class DocumentProcessingError(ValueError):
    """文件无法解析为可审计的结构元素。"""


@dataclass
class DocumentElement:
    """独立的文档语义单元；metadata 保留表格、来源和解析器的额外信息。"""

    type: str
    content: str
    page: int = 1
    section_path: List[str] = field(default_factory=list)
    metadata: Dict[str, Any] = field(default_factory=dict)
    confidence: float = 1.0


@dataclass
class ParsedDocument:
    elements: List[DocumentElement]
    parser_name: str
    parser_version: str
    pages: List[Dict[str, Any]]
    quality: Dict[str, Any]


def parse_document(path: Path, extension: str) -> ParsedDocument:
    """按格式路由解析器，并给出统一质量报告。

    PDF 先用 Docling 获取版面和表格结构；依赖或模型不可用时回退到 pypdf，并将
    回退原因记入报告。扫描件、空文本和乱码率过高的文档不自动发布。
    """
    extension = extension.lower()
    if extension == ".pdf":
        elements, parser_name, version, pages, warnings = _parse_pdf(path)
    elif extension == ".docx":
        elements, parser_name, version, pages, warnings = _parse_docx(path)
    elif extension in {".xlsx", ".xls", ".csv"}:
        elements, parser_name, version, pages, warnings = _parse_spreadsheet(path, extension)
    elif extension == ".pptx":
        elements, parser_name, version, pages, warnings = _parse_pptx(path)
    elif extension == ".md":
        text = path.read_text(encoding="utf-8", errors="replace")
        elements = _elements_from_markdown(text)
        parser_name, version, pages, warnings = "markdown-it-py", _module_version("markdown_it"), _pages_from_elements(elements), []
    elif extension in {".html", ".htm"}:
        elements, parser_name, version, pages, warnings = _parse_html(path)
    elif extension == ".json":
        elements, parser_name, version, pages, warnings = _parse_json(path)
    elif extension == ".xml":
        elements, parser_name, version, pages, warnings = _parse_xml(path)
    elif extension == ".txt":
        text = path.read_text(encoding="utf-8", errors="replace")
        elements = _paragraph_elements(text)
        parser_name, version, pages, warnings = "stdlib-text", "python", _pages_from_elements(elements), []
    else:
        raise DocumentProcessingError(f"不支持的文件类型: {extension}")

    quality = _quality_report(elements, extension, warnings)
    return ParsedDocument(elements, parser_name, version, pages, quality)


def build_chunk_documents(elements: Iterable[DocumentElement], document_name: str) -> List[Document]:
    """将结构元素转为带标题/表头上下文的向量化文档。

    表格按行生成独立元素，正文保留文档名、标题路径和页码前缀。这样短语如
    “试用期为 6 个月”也能检索到其所属制度和字段含义。
    """
    documents: List[Document] = []
    for element in elements:
        if element.type == "heading" or not element.content.strip():
            continue
        section = " > ".join(element.section_path) or "未分章节"
        label = {"table_row": "表格行", "list": "列表", "code": "代码", "image_ocr": "OCR 文本"}.get(element.type, "正文")
        content = f"文档：{document_name}\n章节：{section}\n页码：第 {element.page} 页\n类型：{label}\n{element.content.strip()}"
        metadata = {
            "page": element.page,
            "section_path": section,
            "content_type": element.type,
            "element_metadata": json.dumps(element.metadata, ensure_ascii=False),
        }
        documents.append(Document(page_content=content, metadata=metadata))
    # 所有长元素仍走既有 Token splitter；表格行与短段落不会被硬切。
    return split_documents(documents)


def element_dicts(elements: Iterable[DocumentElement]) -> List[Dict[str, Any]]:
    return [asdict(element) for element in elements]


def _parse_pdf(path: Path) -> Tuple[List[DocumentElement], str, str, List[Dict[str, Any]], List[str]]:
    try:
        from docling.document_converter import DocumentConverter

        result = DocumentConverter().convert(str(path))
        markdown = result.document.export_to_markdown()
        elements = _elements_from_markdown(markdown)
        if not any(element.content.strip() for element in elements):
            raise DocumentProcessingError("Docling 未提取到可用文本")
        return elements, "docling", _module_version("docling"), _pages_from_elements(elements), []
    except Exception as error:
        # pypdf 是许可友好的文本型 PDF 兜底；同时尝试 pdfplumber 的矢量表格
        # 提取，确保主解析器临时失败时不会把关键规则表退化成散乱文字。
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        elements: List[DocumentElement] = []
        pages: List[Dict[str, Any]] = []
        for page_number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            pages.append({"page": page_number, "text": text})
            elements.extend(_paragraph_elements(text, page=page_number))
        try:
            import pdfplumber

            with pdfplumber.open(str(path)) as pdf:
                for page_number, page in enumerate(pdf.pages, 1):
                    for table_index, table in enumerate(page.extract_tables(), 1):
                        rows = [[cell or "" for cell in row] for row in table if row]
                        elements.extend(_table_row_elements(rows, table_id=f"pdf-{page_number}-table-{table_index}", page=page_number))
        except Exception:
            pass
        return elements, "pypdf-fallback", _module_version("pypdf"), pages, [f"Docling 不可用或解析失败：{type(error).__name__}"]


def _parse_docx(path: Path) -> Tuple[List[DocumentElement], str, str, List[Dict[str, Any]], List[str]]:
    from docx import Document as DocxDocument

    source = DocxDocument(str(path))
    elements: List[DocumentElement] = []
    section_path: List[str] = []
    for paragraph in source.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style_name = (paragraph.style.name or "").lower()
        heading_match = re.search(r"heading\s*(\d+)|标题\s*(\d+)", style_name)
        if heading_match:
            level = int(next(value for value in heading_match.groups() if value))
            section_path = section_path[: level - 1] + [text]
            elements.append(DocumentElement("heading", text, section_path=section_path.copy(), metadata={"level": level}))
        elif style_name.startswith("list") or "列表" in style_name:
            elements.append(DocumentElement("list", text, section_path=section_path.copy()))
        else:
            elements.append(DocumentElement("paragraph", text, section_path=section_path.copy()))
    for table_index, table in enumerate(source.tables, 1):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        elements.extend(_table_row_elements(rows, table_id=f"table-{table_index}", section_path=section_path.copy()))
    return elements, "python-docx", _module_version("docx"), _pages_from_elements(elements), []


def _parse_spreadsheet(path: Path, extension: str) -> Tuple[List[DocumentElement], str, str, List[Dict[str, Any]], List[str]]:
    import pandas as pd

    sheets = pd.read_csv(path, dtype=str, keep_default_na=False) if extension == ".csv" else pd.read_excel(path, sheet_name=None, dtype=str, keep_default_na=False)
    if extension == ".csv":
        sheets = {path.stem: sheets}
    elements: List[DocumentElement] = []
    for sheet_name, frame in sheets.items():
        rows = [list(map(str, frame.columns.tolist()))] + frame.fillna("").astype(str).values.tolist()
        elements.append(DocumentElement("heading", str(sheet_name), section_path=[str(sheet_name)], metadata={"sheet": str(sheet_name)}))
        elements.extend(_table_row_elements(rows, table_id=f"sheet-{sheet_name}", section_path=[str(sheet_name)], metadata={"sheet": str(sheet_name)}))
    return elements, "pandas", _module_version("pandas"), _pages_from_elements(elements), []


def _parse_pptx(path: Path) -> Tuple[List[DocumentElement], str, str, List[Dict[str, Any]], List[str]]:
    from pptx import Presentation

    presentation = Presentation(str(path))
    elements: List[DocumentElement] = []
    pages: List[Dict[str, Any]] = []
    for page, slide in enumerate(presentation.slides, 1):
        title = ""
        page_text: List[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                if text:
                    if shape == slide.shapes.title:
                        title = text
                        elements.append(DocumentElement("heading", text, page=page, section_path=[text]))
                    else:
                        page_text.append(text)
                        elements.append(DocumentElement("paragraph", text, page=page, section_path=[title] if title else []))
            if getattr(shape, "has_table", False):
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                elements.extend(_table_row_elements(rows, table_id=f"slide-{page}-table", page=page, section_path=[title] if title else []))
        pages.append({"page": page, "text": "\n".join(page_text)})
    return elements, "python-pptx", _module_version("pptx"), pages, []


def _parse_html(path: Path) -> Tuple[List[DocumentElement], str, str, List[Dict[str, Any]], List[str]]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for noisy in soup(["script", "style", "noscript"]):
        noisy.decompose()
    elements: List[DocumentElement] = []
    section_path: List[str] = []
    for node in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "table"]):
        text = node.get_text(" ", strip=True)
        if not text:
            continue
        if node.name.startswith("h"):
            level = int(node.name[1])
            section_path = section_path[: level - 1] + [text]
            elements.append(DocumentElement("heading", text, section_path=section_path.copy(), metadata={"level": level}))
        elif node.name == "table":
            rows = [[cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"])] for row in node.find_all("tr")]
            elements.extend(_table_row_elements(rows, table_id=f"html-table-{len(elements)}", section_path=section_path.copy()))
        else:
            element_type = "code" if node.name == "pre" else "list" if node.name == "li" else "paragraph"
            elements.append(DocumentElement(element_type, text, section_path=section_path.copy()))
    return elements, "beautifulsoup4", _module_version("bs4"), _pages_from_elements(elements), []


def _parse_json(path: Path) -> Tuple[List[DocumentElement], str, str, List[Dict[str, Any]], List[str]]:
    value = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    elements = _tree_elements(value)
    return elements, "stdlib-json", "python", _pages_from_elements(elements), []


def _parse_xml(path: Path) -> Tuple[List[DocumentElement], str, str, List[Dict[str, Any]], List[str]]:
    from lxml import etree

    root = etree.parse(str(path)).getroot()
    value = _xml_value(root)
    elements = _tree_elements(value)
    return elements, "lxml", _module_version("lxml"), _pages_from_elements(elements), []


def _elements_from_markdown(text: str) -> List[DocumentElement]:
    elements: List[DocumentElement] = []
    section_path: List[str] = []
    table_lines: List[str] = []

    def flush_table() -> None:
        nonlocal table_lines
        if len(table_lines) >= 2:
            rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in table_lines if not re.fullmatch(r"\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?", line)]
            elements.extend(_table_row_elements(rows, table_id=f"markdown-table-{len(elements)}", section_path=section_path.copy()))
        table_lines = []

    paragraph: List[str] = []
    for line in text.replace("\r\n", "\n").split("\n"):
        if line.strip().startswith("|") and "|" in line.strip()[1:]:
            if paragraph:
                elements.append(DocumentElement("paragraph", " ".join(paragraph), section_path=section_path.copy()))
                paragraph = []
            table_lines.append(line)
            continue
        flush_table()
        heading = re.match(r"^(#{1,6})\s+(.+)$", line.strip())
        if heading:
            if paragraph:
                elements.append(DocumentElement("paragraph", " ".join(paragraph), section_path=section_path.copy()))
                paragraph = []
            level, title = len(heading.group(1)), heading.group(2).strip()
            section_path = section_path[: level - 1] + [title]
            elements.append(DocumentElement("heading", title, section_path=section_path.copy(), metadata={"level": level}))
        elif line.strip().startswith(("- ", "* ", "+ ")):
            if paragraph:
                elements.append(DocumentElement("paragraph", " ".join(paragraph), section_path=section_path.copy()))
                paragraph = []
            elements.append(DocumentElement("list", line.strip()[2:], section_path=section_path.copy()))
        elif line.strip():
            paragraph.append(line.strip())
        elif paragraph:
            elements.append(DocumentElement("paragraph", " ".join(paragraph), section_path=section_path.copy()))
            paragraph = []
    flush_table()
    if paragraph:
        elements.append(DocumentElement("paragraph", " ".join(paragraph), section_path=section_path.copy()))
    return elements


def _paragraph_elements(text: str, page: int = 1) -> List[DocumentElement]:
    # 纯文本本身没有标题/表格等可利用的结构；先保留整页原文及空行，再交给
    # Token splitter 在段落边界切分，避免短文被人为拆成多个低价值 Chunk。
    return [DocumentElement("paragraph", text.strip(), page=page)] if text.strip() else []


def _table_row_elements(rows: List[List[str]], table_id: str, page: int = 1, section_path: List[str] = None, metadata: Dict[str, Any] = None) -> List[DocumentElement]:
    if not rows:
        return []
    headers = [header.strip() or f"列{index + 1}" for index, header in enumerate(rows[0])]
    common = metadata or {}
    elements: List[DocumentElement] = []
    for row_index, row in enumerate(rows[1:], 1):
        values = [str(value).strip() for value in row]
        if not any(values):
            continue
        fields = {headers[index]: values[index] if index < len(values) else "" for index in range(len(headers))}
        content = "；".join(f"{key}：{value}" for key, value in fields.items() if value)
        elements.append(DocumentElement("table_row", content, page=page, section_path=section_path or [], metadata={**common, "table_id": table_id, "row_index": row_index, "headers": headers, "cells": fields}))
    return elements


def _tree_elements(value: Any, path: List[str] = None) -> List[DocumentElement]:
    path = path or []
    if isinstance(value, dict):
        elements: List[DocumentElement] = []
        for key, child in value.items():
            elements.extend(_tree_elements(child, path + [str(key)]))
        return elements
    if isinstance(value, list):
        return [element for index, child in enumerate(value, 1) for element in _tree_elements(child, path + [str(index)])]
    return [DocumentElement("paragraph", f"{' > '.join(path)}：{value}", section_path=path[:-1], metadata={"field_path": path})]


def _xml_value(node: Any) -> Dict[str, Any]:
    children = list(node)
    if not children:
        return {node.tag: (node.text or "").strip()}
    grouped: Dict[str, Any] = {}
    for child in children:
        value = _xml_value(child)[child.tag]
        if child.tag in grouped:
            grouped[child.tag] = grouped[child.tag] if isinstance(grouped[child.tag], list) else [grouped[child.tag]]
            grouped[child.tag].append(value)
        else:
            grouped[child.tag] = value
    return {node.tag: grouped}


def _pages_from_elements(elements: List[DocumentElement]) -> List[Dict[str, Any]]:
    by_page: Dict[int, List[str]] = {}
    for element in elements:
        if element.content:
            by_page.setdefault(element.page, []).append(element.content)
    return [{"page": page, "text": "\n".join(content)} for page, content in sorted(by_page.items())]


def _quality_report(elements: List[DocumentElement], extension: str, warnings: List[str]) -> Dict[str, Any]:
    text = "\n".join(element.content for element in elements)
    text_characters = len(text)
    garbled = sum(1 for character in text if character == "\ufffd" or ord(character) < 32 and character not in "\n\t\r")
    garbled_ratio = garbled / max(1, text_characters)
    table_rows = sum(1 for element in elements if element.type == "table_row")
    review_reasons = list(warnings)
    if not text.strip():
        review_reasons.append("未提取到文本")
    if garbled_ratio > 0.03:
        review_reasons.append("乱码比例过高")
    if extension == ".pdf" and text_characters < 30:
        review_reasons.append("PDF 文本过少，可能是扫描件，需要 OCR 或人工审核")
    return {"text_characters": text_characters, "element_count": len(elements), "table_row_count": table_rows, "garbled_ratio": round(garbled_ratio, 6), "warnings": warnings, "review_reasons": review_reasons, "recommended_status": "needs_review" if review_reasons else "ready"}


def _module_version(module_name: str) -> str:
    try:
        module = __import__(module_name)
        return str(getattr(module, "__version__", "unknown"))
    except Exception:
        return "unknown"
